from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    flash,
    jsonify
)

import mysql.connector
from mysql.connector import Error

import os
import re
import json
from functools import wraps

from dotenv import load_dotenv
from huggingface_hub import InferenceClient


# =========================================================
# LOAD ENVIRONMENT VARIABLES
# =========================================================

load_dotenv()


# =========================================================
# FLASK APP
# =========================================================

app = Flask(__name__)

app.secret_key = os.getenv(
    "FLASK_SECRET_KEY",
    "healthcare_secret_key"
)


# =========================================================
# MYSQL CONFIGURATION
# =========================================================

DB_HOST = os.getenv("DB_HOST", "localhost")
DB_USER = os.getenv("DB_USER", "root")
DB_PASSWORD = os.getenv("DB_PASSWORD", "")
DB_NAME = os.getenv("DB_NAME", "healthcare_assistant")


# =========================================================
# HUGGING FACE CONFIGURATION
# =========================================================

HF_TOKEN = os.getenv("HF_TOKEN")

HF_MODEL = os.getenv(
    "HF_MODEL",
    "meta-llama/Llama-3.1-8B-Instruct"
)

if HF_TOKEN:
    hf_client = InferenceClient(token=HF_TOKEN)
else:
    hf_client = None


# =========================================================
# DATABASE CONNECTION
# =========================================================

def get_db_connection():

    try:
        connection = mysql.connector.connect(
            host=DB_HOST,
            user=DB_USER,
            password=DB_PASSWORD,
            database=DB_NAME
        )

        return connection

    except Error as error:

        print("❌ MySQL Connection Error:", error)

        return None


# =========================================================
# DATABASE INITIALIZATION
# =========================================================

def initialize_database():

    connection = get_db_connection()

    if connection is None:
        print("❌ Database initialization skipped.")
        return

    cursor = connection.cursor()

    try:

        # -------------------------------------------------
        # USERS TABLE
        # IMPORTANT LOGIN/REGISTER FIX
        # -------------------------------------------------

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INT AUTO_INCREMENT PRIMARY KEY,
                name VARCHAR(100) NOT NULL,
                email VARCHAR(255) NOT NULL UNIQUE,
                password VARCHAR(255) NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        # -------------------------------------------------
        # PERMANENT PERSONAL MEMORY
        # -------------------------------------------------

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS user_memory (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT NOT NULL,
                memory_key VARCHAR(100) NOT NULL,
                memory_value TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

                UNIQUE KEY unique_user_memory
                (user_id, memory_key),

                INDEX idx_memory_user
                (user_id, id)
            )
        """)

        # -------------------------------------------------
        # AI CHAT MEMORY
        # -------------------------------------------------

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS ai_memory (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT NOT NULL,
                role ENUM('user', 'assistant') NOT NULL,
                message LONGTEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

                INDEX idx_ai_memory_user
                (user_id, id)
            )
        """)

        # -------------------------------------------------
        # MEDICAL REPORTS
        # -------------------------------------------------

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS medical_reports (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT NOT NULL,
                report_name VARCHAR(255),
                report_text LONGTEXT,
                ai_analysis LONGTEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

                INDEX idx_medical_user
                (user_id, created_at)
            )
        """)

        # -------------------------------------------------
        # BLOOD ANALYSIS
        # -------------------------------------------------

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS blood_analysis (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT NOT NULL,
                hemoglobin VARCHAR(100),
                blood_glucose VARCHAR(100),
                wbc VARCHAR(100),
                platelets VARCHAR(100),
                rbc VARCHAR(100),
                cholesterol VARCHAR(100),
                blood_pressure VARCHAR(100),
                other_values TEXT,
                ai_analysis LONGTEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

                INDEX idx_blood_user
                (user_id, created_at)
            )
        """)

        # -------------------------------------------------
        # SYMPTOM CHECKS
        # -------------------------------------------------

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS symptom_checks (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT NOT NULL,
                symptoms TEXT NOT NULL,
                duration VARCHAR(100),
                severity VARCHAR(50),
                ai_analysis LONGTEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

                INDEX idx_symptom_user
                (user_id, created_at)
            )
        """)

        # -------------------------------------------------
        # HEALTH INSIGHTS
        # -------------------------------------------------

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS health_insights (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT NOT NULL,
                insight LONGTEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,

                INDEX idx_insight_user
                (user_id, created_at)
            )
        """)

        connection.commit()

        print("✅ All application tables initialized successfully.")

    except Error as error:

        print("❌ Database initialization error:", error)

        connection.rollback()

    finally:

        cursor.close()
        connection.close()


# =========================================================
# LOGIN REQUIRED DECORATOR
# =========================================================

def login_required(function):

    @wraps(function)
    def decorated_function(*args, **kwargs):

        if "user_id" not in session:

            api_paths = (
                "/ask-ai",
                "/analyze-",
                "/check-",
                "/generate-",
                "/medical-report-history",
                "/blood-analysis-history",
                "/symptom-history",
                "/health-insights-history",
                "/clear-",
                "/memory",
                "/chat-history"
            )

            if request.path.startswith(api_paths):

                return jsonify({
                    "success": False,
                    "error": "Please login first."
                }), 401

            return redirect(url_for("login"))

        return function(*args, **kwargs)

    return decorated_function


# =========================================================
# HOME
# =========================================================

@app.route("/")
def home():

    if "user_id" in session:
        return redirect(url_for("dashboard"))

    return redirect(url_for("login"))


# =========================================================
# REGISTER
# =========================================================

@app.route("/register", methods=["GET", "POST"])
def register():

    if request.method == "POST":

        name = request.form.get(
            "name",
            ""
        ).strip()

        email = request.form.get(
            "email",
            ""
        ).strip().lower()

        password = request.form.get(
            "password",
            ""
        ).strip()

        if not name or not email or not password:

            flash(
                "All fields are required.",
                "error"
            )

            return redirect(url_for("register"))

        connection = get_db_connection()

        if connection is None:

            flash(
                "Database connection failed.",
                "error"
            )

            return redirect(url_for("register"))

        cursor = connection.cursor()

        try:

            cursor.execute(
                """
                SELECT id
                FROM users
                WHERE email = %s
                """,
                (email,)
            )

            existing_user = cursor.fetchone()

            if existing_user:

                flash(
                    "Email already registered. Please login.",
                    "error"
                )

                return redirect(url_for("login"))

            cursor.execute(
                """
                INSERT INTO users
                (name, email, password)
                VALUES (%s, %s, %s)
                """,
                (
                    name,
                    email,
                    password
                )
            )

            connection.commit()

            print(
                f"✅ New user registered: {email}"
            )

            flash(
                "Registration successful. Please login.",
                "success"
            )

            return redirect(url_for("login"))

        except Error as error:

            print(
                "❌ Registration Error:",
                error
            )

            connection.rollback()

            flash(
                "Registration failed: " + str(error),
                "error"
            )

        finally:

            cursor.close()
            connection.close()

    return render_template("register.html")


# =========================================================
# LOGIN
# =========================================================

@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        email = request.form.get(
            "email",
            ""
        ).strip().lower()

        password = request.form.get(
            "password",
            ""
        ).strip()

        # -------------------------------------------------
        # BASIC VALIDATION
        # -------------------------------------------------

        if not email or not password:

            return render_template(
                "login.html",
                error="Please enter both email and password."
            )

        connection = get_db_connection()

        if connection is None:

            return render_template(
                "login.html",
                error="Database connection failed. Please check MySQL."
            )

        cursor = connection.cursor(dictionary=True)

        try:

            # -------------------------------------------------
            # FIRST FIND USER BY EMAIL
            # -------------------------------------------------

            cursor.execute(
                """
                SELECT *
                FROM users
                WHERE LOWER(email) = %s
                LIMIT 1
                """,
                (email,)
            )

            user = cursor.fetchone()

            # -------------------------------------------------
            # EMAIL NOT FOUND
            # -------------------------------------------------

            if not user:

                print(
                    f"❌ Login failed - email not found: {email}"
                )

                return render_template(
                    "login.html",
                    error=(
                        "No account found with this email. "
                        "Please register first."
                    )
                )

            # -------------------------------------------------
            # PASSWORD CHECK
            # -------------------------------------------------

            if str(user.get("password", "")) != password:

                print(
                    f"❌ Login failed - wrong password: {email}"
                )

                return render_template(
                    "login.html",
                    error="Incorrect password. Please try again."
                )

            # -------------------------------------------------
            # SUCCESSFUL LOGIN
            # -------------------------------------------------

            session.clear()

            session["user_id"] = user["id"]

            session["user_name"] = (
                user.get("name")
                or user.get("username")
                or "User"
            )

            session["user_email"] = (
                user.get("email")
                or email
            )

            print(
                f"✅ Login successful: {email}"
            )

            return redirect(
                url_for("dashboard")
            )

        except Error as error:

            print(
                "❌ Login Database Error:",
                error
            )

            return render_template(
                "login.html",
                error="Login failed: " + str(error)
            )

        except Exception as error:

            print(
                "❌ Login Error:",
                repr(error)
            )

            return render_template(
                "login.html",
                error="Something went wrong during login."
            )

        finally:

            cursor.close()
            connection.close()

    return render_template("login.html")


# =========================================================
# DASHBOARD
# =========================================================

@app.route("/dashboard")
@login_required
def dashboard():

    return render_template(
        "dashboard.html"
    )


# =========================================================
# AI ASSISTANT PAGE
# =========================================================

@app.route("/ai-assistant")
@login_required
def ai_assistant():

    return render_template(
        "ai_assistant.html"
    )


# =========================================================
# LOGOUT
# =========================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect(
        url_for("login")
    )


# =========================================================
# HUGGING FACE AI
# =========================================================

def ask_huggingface(
    system_prompt,
    user_prompt,
    max_tokens=700,
    temperature=0.4
):

    if not hf_client:

        raise Exception(
            "Hugging Face token is not configured. "
            "Please add HF_TOKEN to your .env file."
        )

    response = hf_client.chat.completions.create(

        model=HF_MODEL,

        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_prompt
            }
        ],

        max_tokens=max_tokens,
        temperature=temperature,
        top_p=0.9
    )

    if not response:

        raise Exception(
            "Empty response received from Hugging Face."
        )

    answer = (
        response
        .choices[0]
        .message
        .content
    )

    if not answer:

        raise Exception(
            "AI returned an empty response."
        )

    return answer.strip()


# =========================================================
# HEALTHCARE AI SYSTEM PROMPT
# =========================================================

HEALTH_SYSTEM_PROMPT = """

You are HealthCare AI, a healthcare education assistant.

Your purpose is to provide clear, simple and responsible
general health information.

IMPORTANT RULES:

1. Do not provide a confirmed medical diagnosis.

2. Do not prescribe prescription medicines.

3. Do not invent medical test values.

4. Use only values supplied by the user.

5. Explain medical terms in simple language.

6. Laboratory reference ranges can vary by laboratory,
age, sex, pregnancy status, medical history and other factors.

7. Do not assume that one normal test proves overall health.

8. Do not claim that one normal FSH value automatically proves
fertility or a healthy menstrual cycle.

9. Hormone interpretation may depend on menstrual-cycle timing,
symptoms and other laboratory tests.

10. If symptoms may indicate an emergency, clearly recommend
prompt professional medical evaluation.

11. Never create facts that are not present in supplied data.

12. Use headings and bullet points where helpful.

13. Keep responses concise but useful.

14. When personal memory is provided, use it naturally.

15. Never reveal database details, internal memory keys,
user IDs or private system information.

16. If the user asks "What is my name?", use saved profile
or permanent memory when available.

17. New conversations are fresh conversations, but saved
long-term personal facts may still be used.

18. Normal non-medical questions can be answered naturally.

19. Do not repeat the same information unnecessarily.

20. If the user provides previous health-check information,
do not invent additional symptoms or results.

21. For health concerns, clearly separate:
    - possible/common explanations
    - self-care
    - warning signs
    - when professional care is appropriate.

22. When answering questions about personal facts,
use only information explicitly available in memory.
Never guess a personal fact.

"""


# =========================================================
# USER PROFILE
# =========================================================

def get_user_profile(user_id):

    connection = get_db_connection()

    if connection is None:
        return {}

    cursor = connection.cursor(dictionary=True)

    try:

        cursor.execute(
            """
            SELECT
                id,
                name,
                email
            FROM users
            WHERE id = %s
            """,
            (user_id,)
        )

        return cursor.fetchone() or {}

    except Error as error:

        print(
            "User profile error:",
            error
        )

        return {}

    finally:

        cursor.close()
        connection.close()


# =========================================================
# GET PERMANENT MEMORY
# =========================================================

def get_user_memory(
    user_id,
    limit=100
):

    connection = get_db_connection()

    if connection is None:
        return []

    cursor = connection.cursor(dictionary=True)

    try:

        limit = max(
            1,
            min(int(limit), 100)
        )

        cursor.execute(
            f"""
            SELECT
                id,
                memory_key,
                memory_value,
                created_at
            FROM user_memory
            WHERE user_id = %s
            ORDER BY id DESC
            LIMIT {limit}
            """,
            (user_id,)
        )

        return cursor.fetchall()

    except Error as error:

        print(
            "User memory read error:",
            error
        )

        return []

    finally:

        cursor.close()
        connection.close()


# =========================================================
# SAVE / UPDATE PERMANENT MEMORY
# =========================================================

def save_user_memory(
    user_id,
    memory_key,
    memory_value
):

    memory_key = str(
        memory_key
    ).strip().lower()

    memory_value = str(
        memory_value
    ).strip()

    if not memory_key or not memory_value:
        return False

    if len(memory_key) > 100:
        return False

    if len(memory_value) > 1000:
        memory_value = memory_value[:1000]

    connection = get_db_connection()

    if connection is None:
        return False

    cursor = connection.cursor()

    try:

        cursor.execute(
            """
            INSERT INTO user_memory
            (
                user_id,
                memory_key,
                memory_value
            )
            VALUES (%s, %s, %s)

            ON DUPLICATE KEY UPDATE
                memory_value = VALUES(memory_value)
            """,
            (
                user_id,
                memory_key,
                memory_value
            )
        )

        connection.commit()

        return True

    except Error as error:

        print(
            "Memory save error:",
            error
        )

        connection.rollback()

        return False

    finally:

        cursor.close()
        connection.close()


# =========================================================
# BASIC MEMORY EXTRACTION
# =========================================================

def extract_basic_personal_memory(message):

    memories = []

    text = str(message).strip()

    if not text:
        return memories

    # NAME
    name_patterns = [

        r"^\s*my name is\s+([A-Za-z][A-Za-z .'-]{1,49})\s*[.!?]?\s*$",

        r"^\s*my name's\s+([A-Za-z][A-Za-z .'-]{1,49})\s*[.!?]?\s*$",

        r"^\s*call me\s+([A-Za-z][A-Za-z .'-]{1,49})\s*[.!?]?\s*$",

        r"^\s*i am\s+([A-Za-z][A-Za-z .'-]{1,49})\s*[.!?]?\s*$"
    ]

    for pattern in name_patterns:

        match = re.search(
            pattern,
            text,
            re.IGNORECASE
        )

        if match:

            name = match.group(1).strip()

            name = re.split(
                r"[.!?,]",
                name
            )[0].strip()

            if (
                1 <= len(name.split()) <= 4
                and len(name) <= 50
            ):

                memories.append(
                    ("name", name)
                )

                break

    # FAVORITE
    favorite_patterns = [

        r"^\s*my\s+(favorite|favourite)\s+([A-Za-z][A-Za-z ]{1,40})\s+is\s+(.+?)\s*[.!?]?\s*$",

        r"^\s*(.+?)\s+is\s+my\s+(favorite|favourite)\s+([A-Za-z][A-Za-z ]{1,40})\s*[.!?]?\s*$"
    ]

    match = re.search(
        favorite_patterns[0],
        text,
        re.IGNORECASE
    )

    if match:

        category = match.group(2).strip()
        value = match.group(3).strip()

        value = re.split(
            r"[.!?]",
            value
        )[0].strip()

        category = re.sub(
            r"\s+",
            "_",
            category.lower()
        )

        if 1 <= len(value) <= 100:

            memories.append(
                (
                    f"favorite_{category}",
                    value
                )
            )

    else:

        match = re.search(
            favorite_patterns[1],
            text,
            re.IGNORECASE
        )

        if match:

            value = match.group(1).strip()
            category = match.group(3).strip()

            value = re.split(
                r"[.!?]",
                value
            )[0].strip()

            category = re.sub(
                r"\s+",
                "_",
                category.lower()
            )

            if 1 <= len(value) <= 100:

                memories.append(
                    (
                        f"favorite_{category}",
                        value
                    )
                )

    # LIKE
    like_patterns = [

        r"^\s*i\s+(?:really\s+)?like\s+(.+?)\s*[.!?]?\s*$",

        r"^\s*i\s+(?:really\s+)?love\s+(.+?)\s*[.!?]?\s*$",

        r"^\s*i\s+(?:really\s+)?enjoy\s+(.+?)\s*[.!?]?\s*$",

        r"^\s*i\s+prefer\s+(.+?)\s*[.!?]?\s*$"
    ]

    for pattern in like_patterns:

        match = re.search(
            pattern,
            text,
            re.IGNORECASE
        )

        if match:

            value = match.group(1).strip()

            value = re.split(
                r"[.!?]",
                value
            )[0].strip()

            if 2 <= len(value) <= 100:

                memories.append(
                    (
                        "preference",
                        value
                    )
                )

            break

    # DISLIKE
    dislike_patterns = [

        r"^\s*i\s+(?:really\s+)?(?:don't|do not)\s+like\s+(.+?)\s*[.!?]?\s*$",

        r"^\s*i\s+(?:really\s+)?dislike\s+(.+?)\s*[.!?]?\s*$",

        r"^\s*i\s+hate\s+(.+?)\s*[.!?]?\s*$"
    ]

    for pattern in dislike_patterns:

        match = re.search(
            pattern,
            text,
            re.IGNORECASE
        )

        if match:

            value = match.group(1).strip()

            value = re.split(
                r"[.!?]",
                value
            )[0].strip()

            if 2 <= len(value) <= 100:

                memories.append(
                    (
                        "dislike",
                        value
                    )
                )

            break

    return memories


# =========================================================
# AI MEMORY EXTRACTION
# =========================================================

MEMORY_EXTRACTION_PROMPT = """

You are a personal-memory extraction component.

Read the user's message and identify ONLY stable,
useful personal facts or preferences that may help
the assistant in future conversations.

Do NOT save:
- temporary moods
- one-time events
- medical diagnoses
- passwords
- OTPs
- financial information
- addresses
- phone numbers
- email addresses
- highly sensitive information
- private secrets
- questions
- requests
- instructions to the assistant

Good examples:

"My favorite color is blue"
-> favorite_color = blue

"I like Java"
-> preference = Java

"I am preparing for placements"
-> goal = preparing for placements

"I am a computer science student"
-> education = computer science student

"I prefer studying in the morning"
-> preference_study_time = morning

Important:
- Do not infer facts.
- Do not guess.
- Only extract what is explicitly stated.
- If nothing should be remembered, return [].

Return ONLY valid JSON.

Format:

[
  {
    "key": "short_key",
    "value": "explicit_value"
  }
]

"""


def extract_ai_personal_memory(message):

    if not hf_client:
        return []

    text = str(message).strip()

    if not text:
        return []

    text = text[:3000]

    try:

        response = hf_client.chat.completions.create(

            model=HF_MODEL,

            messages=[
                {
                    "role": "system",
                    "content": MEMORY_EXTRACTION_PROMPT
                },
                {
                    "role": "user",
                    "content": text
                }
            ],

            max_tokens=300,
            temperature=0.0,
            top_p=1.0
        )

        raw = (
            response
            .choices[0]
            .message
            .content
        )

        if not raw:
            return []

        raw = raw.strip()

        raw = re.sub(
            r"^```(?:json)?",
            "",
            raw,
            flags=re.IGNORECASE
        )

        raw = re.sub(
            r"```$",
            "",
            raw
        )

        raw = raw.strip()

        data = json.loads(raw)

        if not isinstance(data, list):
            return []

        memories = []

        for item in data:

            if not isinstance(item, dict):
                continue

            key = str(
                item.get("key", "")
            ).strip()

            value = str(
                item.get("value", "")
            ).strip()

            if not key or not value:
                continue

            key = re.sub(
                r"[^a-zA-Z0-9_]",
                "_",
                key.lower()
            )

            key = re.sub(
                r"_+",
                "_",
                key
            ).strip("_")

            if not key:
                continue

            key = key[:100]
            value = value[:500]

            memories.append(
                (
                    key,
                    value
                )
            )

        return memories[:10]

    except Exception as error:

        print(
            "AI memory extraction skipped:",
            repr(error)
        )

        return []


# =========================================================
# REMEMBER FROM MESSAGE
# =========================================================

def remember_from_message(
    user_id,
    message
):

    basic_memories = extract_basic_personal_memory(
        message
    )

    for key, value in basic_memories:

        save_user_memory(
            user_id,
            key,
            value
        )

    ai_memories = extract_ai_personal_memory(
        message
    )

    for key, value in ai_memories:

        if key == "name":
            continue

        save_user_memory(
            user_id,
            key,
            value
        )


# =========================================================
# PERSONAL CONTEXT
# =========================================================

def build_personal_context(user_id):

    profile = get_user_profile(user_id)

    memories = get_user_memory(
        user_id,
        100
    )

    parts = []

    if profile.get("name"):

        parts.append(
            f"Name: {profile['name']}"
        )

    for memory in reversed(memories):

        key = memory.get("memory_key")
        value = memory.get("memory_value")

        if key and value:

            parts.append(
                f"{key}: {value}"
            )

    if not parts:

        return "No saved personal information."

    return "\n".join(parts)


# =========================================================
# CHAT HISTORY
# =========================================================

def get_recent_chat_history(
    user_id,
    limit=12
):

    connection = get_db_connection()

    if connection is None:
        return []

    cursor = connection.cursor(
        dictionary=True
    )

    try:

        limit = max(
            1,
            min(int(limit), 50)
        )

        cursor.execute(
            f"""
            SELECT
                id,
                role,
                message,
                created_at
            FROM ai_memory
            WHERE user_id = %s
            ORDER BY id DESC
            LIMIT {limit}
            """,
            (user_id,)
        )

        rows = cursor.fetchall()

        rows.reverse()

        return rows

    except Error as error:

        print(
            "Chat history error:",
            error
        )

        return []

    finally:

        cursor.close()
        connection.close()


# =========================================================
# SAVE CHAT
# =========================================================

def save_ai_memory(
    user_id,
    role,
    message
):

    if role not in (
        "user",
        "assistant"
    ):
        return False

    if not message:
        return False

    connection = get_db_connection()

    if connection is None:
        return False

    cursor = connection.cursor()

    try:

        cursor.execute(
            """
            INSERT INTO ai_memory
            (
                user_id,
                role,
                message
            )
            VALUES (%s, %s, %s)
            """,
            (
                user_id,
                role,
                message
            )
        )

        connection.commit()

        return True

    except Error as error:

        print(
            "Chat memory save error:",
            error
        )

        connection.rollback()

        return False

    finally:

        cursor.close()
        connection.close()


# =========================================================
# CHAT HISTORY API
# =========================================================

@app.route("/chat-history", methods=["GET"])
@login_required
def chat_history():

    try:

        history = get_recent_chat_history(
            session["user_id"],
            100
        )

        for item in history:

            if item.get("created_at"):

                item["created_at"] = (
                    item["created_at"]
                    .strftime(
                        "%d %b %Y, %I:%M %p"
                    )
                )

        return jsonify({
            "success": True,
            "history": history
        })

    except Exception as error:

        print(
            "Chat history API error:",
            repr(error)
        )

        return jsonify({
            "success": False,
            "error": "Unable to load chat history."
        }), 500


# =========================================================
# ASK AI
# =========================================================

@app.route("/ask-ai", methods=["POST"])
@login_required
def ask_ai():

    try:

        data = request.get_json(
            silent=True
        ) or {}

        message = str(
            data.get(
                "message",
                ""
            )
        ).strip()

        if not message:

            return jsonify({
                "success": False,
                "error": "Please enter a message."
            }), 400

        user_id = session["user_id"]

        remember_from_message(
            user_id,
            message
        )

        history = get_recent_chat_history(
            user_id,
            10
        )

        personal_context = build_personal_context(
            user_id
        )

        messages = [

            {
                "role": "system",
                "content": HEALTH_SYSTEM_PROMPT
            },

            {
                "role": "system",
                "content":
                    "LONG-TERM PERSONAL MEMORY:\n"
                    + personal_context
            }
        ]

        for item in history:

            role = item.get("role")
            content = item.get("message")

            if role in (
                "user",
                "assistant"
            ) and content:

                messages.append({
                    "role": role,
                    "content": content
                })

        messages.append({
            "role": "user",
            "content": message
        })

        if not hf_client:

            return jsonify({
                "success": False,
                "error":
                    "Hugging Face token is not configured."
            }), 500

        response = hf_client.chat.completions.create(

            model=HF_MODEL,
            messages=messages,
            max_tokens=700,
            temperature=0.4,
            top_p=0.9
        )

        answer = (
            response
            .choices[0]
            .message
            .content
        )

        if not answer:
            raise Exception(
                "AI returned an empty answer."
            )

        answer = answer.strip()

        save_ai_memory(
            user_id,
            "user",
            message
        )

        save_ai_memory(
            user_id,
            "assistant",
            answer
        )

        return jsonify({
            "success": True,
            "answer": answer
        })

    except Exception as error:

        print(
            "❌ /ask-ai Error:",
            repr(error)
        )

        return jsonify({
            "success": False,
            "error":
                "AI service error: " + str(error)
        }), 500


# =========================================================
# NEW CHAT
# =========================================================

@app.route("/clear-ai-memory", methods=["POST"])
@login_required
def clear_ai_memory():

    user_id = session["user_id"]

    connection = get_db_connection()

    if connection is None:

        return jsonify({
            "success": False,
            "error":
                "Database connection failed."
        }), 500

    cursor = connection.cursor()

    try:

        cursor.execute(
            """
            DELETE FROM ai_memory
            WHERE user_id = %s
            """,
            (user_id,)
        )

        connection.commit()

        return jsonify({
            "success": True,
            "message":
                "New chat started successfully. "
                "Permanent memory was preserved."
        })

    except Error as error:

        print(
            "Clear chat error:",
            error
        )

        connection.rollback()

        return jsonify({
            "success": False,
            "error":
                "Unable to start a new chat."
        }), 500

    finally:

        cursor.close()
        connection.close()


# =========================================================
# VIEW PERMANENT MEMORY
# =========================================================

@app.route("/memory", methods=["GET"])
@login_required
def memory():

    memories = get_user_memory(
        session["user_id"],
        100
    )

    for item in memories:

        if item.get("created_at"):

            item["created_at"] = (
                item["created_at"]
                .strftime(
                    "%d %b %Y, %I:%M %p"
                )
            )

    return jsonify({
        "success": True,
        "memories": memories
    })


# =========================================================
# DELETE ONE MEMORY
# =========================================================

@app.route("/memory/<int:memory_id>", methods=["DELETE"])
@login_required
def delete_memory(memory_id):

    connection = get_db_connection()

    if connection is None:

        return jsonify({
            "success": False,
            "error":
                "Database connection failed."
        }), 500

    cursor = connection.cursor()

    try:

        cursor.execute(
            """
            DELETE FROM user_memory
            WHERE id = %s
            AND user_id = %s
            """,
            (
                memory_id,
                session["user_id"]
            )
        )

        connection.commit()

        return jsonify({
            "success": True,
            "message":
                "Memory deleted successfully."
        })

    except Error as error:

        print(
            "Memory delete error:",
            error
        )

        connection.rollback()

        return jsonify({
            "success": False,
            "error":
                "Unable to delete memory."
        }), 500

    finally:

        cursor.close()
        connection.close()


# =========================================================
# CLEAR ALL PERMANENT MEMORY
# =========================================================

@app.route(
    "/memory/clear",
    methods=["DELETE", "POST"]
)
@login_required
def clear_all_memory():

    connection = get_db_connection()

    if connection is None:

        return jsonify({
            "success": False,
            "error":
                "Database connection failed."
        }), 500

    cursor = connection.cursor()

    try:

        cursor.execute(
            """
            DELETE FROM user_memory
            WHERE user_id = %s
            """,
            (session["user_id"],)
        )

        connection.commit()

        return jsonify({
            "success": True,
            "message":
                "All permanent personal memory deleted."
        })

    except Error as error:

        print(
            "Clear permanent memory error:",
            error
        )

        connection.rollback()

        return jsonify({
            "success": False,
            "error":
                "Unable to delete permanent memory."
        }), 500

    finally:

        cursor.close()
        connection.close()


# =========================================================
# MEDICAL REPORT EXTRACTION + OCR
# =========================================================

ALLOWED_MEDICAL_REPORT_EXTENSIONS = {
    "pdf",
    "txt",
    "docx",
    "png",
    "jpg",
    "jpeg",
    "webp"
}


def _get_file_extension(filename):

    filename = filename or ""

    if "." not in filename:
        return ""

    return filename.rsplit(
        ".",
        1
    )[-1].lower().strip()


def _extract_image_text(image):

    try:
        import pytesseract
    except ImportError:

        raise Exception(
            "Image OCR is not installed. "
            "Run: pip install pytesseract pillow"
        )

    try:

        image = image.convert("RGB")

        text = pytesseract.image_to_string(
            image
        )

        return text or ""

    except Exception as error:

        raise Exception(
            "Unable to read text from the uploaded image. "
            "Make sure Tesseract OCR is installed on Windows. "
            f"Details: {error}"
        )


def _ocr_pdf(file_bytes):

    try:
        import fitz
    except ImportError:

        raise Exception(
            "Scanned PDF OCR requires PyMuPDF. "
            "Run: pip install pymupdf pytesseract pillow"
        )

    try:
        import pytesseract
    except ImportError:

        raise Exception(
            "Scanned PDF OCR requires pytesseract. "
            "Run: pip install pytesseract pillow"
        )

    try:
        from PIL import Image
    except ImportError:

        raise Exception(
            "Scanned PDF OCR requires Pillow. "
            "Run: pip install pillow"
        )

    document = fitz.open(
        stream=file_bytes,
        filetype="pdf"
    )

    pages = []

    try:

        max_pages = min(
            len(document),
            20
        )

        for page_number in range(max_pages):

            page = document.load_page(
                page_number
            )

            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(2, 2),
                alpha=False
            )

            image = Image.frombytes(
                "RGB",
                [
                    pixmap.width,
                    pixmap.height
                ],
                pixmap.samples
            )

            text = pytesseract.image_to_string(
                image
            )

            if text and text.strip():

                pages.append(
                    f"[Page {page_number + 1}]\n"
                    f"{text.strip()}"
                )

    finally:

        document.close()

    return "\n\n".join(pages)


def extract_report_text(file):

    filename = file.filename or ""

    extension = _get_file_extension(
        filename
    )

    if extension not in ALLOWED_MEDICAL_REPORT_EXTENSIONS:

        raise Exception(
            "Unsupported report format. Supported formats: "
            "PDF, TXT, DOCX, PNG, JPG/JPEG and WEBP."
        )

    # TXT
    if extension == "txt":

        raw = file.read()

        return raw.decode(
            "utf-8",
            errors="ignore"
        )

    # DOCX
    if extension == "docx":

        try:
            from docx import Document
        except ImportError:

            raise Exception(
                "DOCX support is not installed. "
                "Run: pip install python-docx"
            )

        try:

            document = Document(file)

            paragraphs = [
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            table_rows = []

            for table in document.tables:

                for row in table.rows:

                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    if any(cells):

                        table_rows.append(
                            " | ".join(cells)
                        )

            return "\n".join(
                paragraphs + table_rows
            )

        except Exception as error:

            raise Exception(
                "Unable to read the DOCX medical report. "
                f"Details: {error}"
            )

    # IMAGE
    if extension in {
        "png",
        "jpg",
        "jpeg",
        "webp"
    }:

        try:
            from PIL import Image
        except ImportError:

            raise Exception(
                "Image support is not installed. "
                "Run: pip install pillow"
            )

        try:

            image = Image.open(file)

            return _extract_image_text(image)

        except Exception as error:

            if str(error).startswith(
                "Image OCR is not installed"
            ):
                raise

            raise Exception(
                "Unable to process the uploaded "
                "medical-report image. "
                f"Details: {error}"
            )

    # PDF
    if extension == "pdf":

        file_bytes = file.read()

        try:
            from pypdf import PdfReader
        except ImportError:

            raise Exception(
                "PDF support is not installed. "
                "Run: pip install pypdf"
            )

        try:

            from io import BytesIO

            reader = PdfReader(
                BytesIO(file_bytes)
            )

            pages = []

            for page_number, page in enumerate(
                reader.pages
            ):

                text = page.extract_text()

                if text and text.strip():

                    pages.append(
                        f"[Page {page_number + 1}]\n"
                        f"{text.strip()}"
                    )

            normal_text = "\n\n".join(
                pages
            ).strip()

            if len(normal_text) >= 50:

                return normal_text

            ocr_text = _ocr_pdf(
                file_bytes
            ).strip()

            if ocr_text:
                return ocr_text

            return normal_text

        except Exception as error:

            if any(
                phrase in str(error)
                for phrase in (
                    "Scanned PDF OCR requires",
                    "requires pytesseract",
                    "requires Pillow"
                )
            ):
                raise

            raise Exception(
                "Unable to read the PDF medical report. "
                f"Details: {error}"
            )

    raise Exception(
        "Unable to extract medical report text."
    )


# =========================================================
# UNIVERSAL MEDICAL REPORT ANALYZER
# =========================================================

@app.route(
    "/analyze-medical-report",
    methods=["POST"]
)
@login_required
def analyze_medical_report():

    try:

        if "report" not in request.files:

            return jsonify({
                "success": False,
                "error":
                    "No medical report file was uploaded."
            }), 400

        file = request.files["report"]

        if not file or not file.filename:

            return jsonify({
                "success": False,
                "error":
                    "Please select a medical report."
            }), 400

        extension = _get_file_extension(
            file.filename
        )

        if extension not in ALLOWED_MEDICAL_REPORT_EXTENSIONS:

            return jsonify({
                "success": False,
                "error": (
                    "Unsupported file type. Please upload "
                    "PDF, TXT, DOCX, PNG, JPG/JPEG or WEBP."
                )
            }), 400

        report_text = extract_report_text(
            file
        ).strip()

        if not report_text:

            return jsonify({
                "success": False,
                "error": (
                    "Could not extract readable text from this report. "
                    "If it is a scanned/image report, make sure the text "
                    "is clear and OCR is installed."
                )
            }), 400

        report_text = report_text[:30000]

        prompt = f"""

You are analyzing a medical report for educational
understanding only.

The user may upload ANY common type of medical report,
including laboratory reports, CBC, LFT, KFT/RFT,
lipid profile, thyroid/hormone tests, urine tests,
HbA1c, vitamin tests, pathology reports, ECG reports,
radiology reports, ultrasound, X-ray, CT, MRI reports,
or other clinical reports.

FIRST identify the report type from the supplied text
when possible.

Do not guess a report type if the evidence is insufficient.

FOR NUMERICAL LABORATORY RESULTS:

- Read the test name, result, unit and reference range
  when present.
- Compare the result ONLY with the reference range
  printed in this report.
- Clearly label each result as NORMAL, HIGH, LOW,
  or ABNORMAL when the supplied information supports
  that classification.
- If no reference range is provided, say:
  "Reference range not provided"
  instead of inventing one.
- Do not use a generic reference range in place of
  the report's own range.

FOR NON-NUMERICAL REPORTS SUCH AS ECG, X-RAY, CT,
MRI, ULTRASOUND, PATHOLOGY OR OTHER REPORTS:

- Summarize the written FINDINGS and IMPRESSION
  in simple language.
- Point out findings that the report itself describes
  as abnormal, significant, suspicious, or requiring
  follow-up.
- Do not diagnose a condition that is not explicitly
  established in the report.
- Do not claim to visually interpret an image unless
  reliable text findings are actually supplied.

RETURN THE RESPONSE USING THIS STRUCTURE:

## 1. Report Type

State the detected report/test type.

## 2. Overall Summary

Give 2-4 simple sentences about the main result
of the report.

## 3. Results / Findings

For laboratory values, use:

Test | Result | Unit | Reference Range | Status

For non-laboratory reports, use:

Finding | What the report says | Simple meaning

## 4. What Is Normal

List important normal results/findings supported
by the report.

## 5. What Needs Attention

List HIGH, LOW, ABNORMAL, significant or follow-up
findings.

For each one explain its basic meaning in simple language.

## 6. What You Can Discuss With a Doctor

Give practical questions/topics to discuss with a
qualified healthcare professional.

Do not prescribe medicines or give a personalized
treatment plan.

## 7. General Guidance

Give safe, general health guidance only when appropriate.

Do not tell the user to start or stop prescription medicines.

## 8. When to Seek Prompt Medical Care

Mention urgent warning signs only when relevant to
the supplied report or findings.

Do not create symptoms that are not supplied.

## 9. Important Limitation

Clearly state that this is an educational explanation
and not a diagnosis.

Interpretation can depend on symptoms, medical history,
age, sex, medications, pregnancy status and clinician
assessment.

STRICT RULES:

- Use ONLY information contained in the supplied report.
- Do NOT invent test values.
- Do NOT invent units.
- Do NOT invent reference ranges.
- Do NOT invent findings.
- Do NOT invent symptoms.
- Do NOT diagnose.
- Do NOT claim that one report confirms or rules out
  a disease.
- Do NOT claim that one hormone test proves fertility
  or infertility.
- Do NOT recommend prescription medicines.
- Preserve uncertainty when the report is unclear.
- Keep the language simple.
- Do not treat an OCR mistake as a confirmed medical value.
- If a value looks unclear, say it should be verified
  against the original report.

MEDICAL REPORT:

{report_text}

"""

        analysis = ask_huggingface(
            HEALTH_SYSTEM_PROMPT,
            prompt,
            max_tokens=1400,
            temperature=0.2
        )

        connection = get_db_connection()

        if connection is None:

            return jsonify({
                "success": False,
                "error":
                    "Database connection failed while saving report."
            }), 500

        cursor = connection.cursor()

        try:

            cursor.execute(
                """
                INSERT INTO medical_reports
                (
                    user_id,
                    report_name,
                    report_text,
                    ai_analysis
                )
                VALUES (%s, %s, %s, %s)
                """,
                (
                    session["user_id"],
                    file.filename,
                    report_text,
                    analysis
                )
            )

            connection.commit()

        finally:

            cursor.close()
            connection.close()

        return jsonify({

            "success": True,

            "filename": file.filename,

            "file_type": extension,

            "analysis": analysis

        })

    except Exception as error:

        print(
            "❌ Medical Report Error:",
            repr(error)
        )

        return jsonify({

            "success": False,

            "error":
                "Unable to analyze medical report: "
                + str(error)

        }), 500


# =========================================================
# MEDICAL REPORT HISTORY
# =========================================================

@app.route(
    "/medical-report-history",
    methods=["GET"]
)
@login_required
def medical_report_history():

    connection = get_db_connection()

    if connection is None:

        return jsonify({
            "success": False,
            "error":
                "Database connection failed."
        }), 500

    cursor = connection.cursor(
        dictionary=True
    )

    try:

        cursor.execute(
            """
            SELECT
                id,
                report_name,
                report_text,
                ai_analysis,
                created_at
            FROM medical_reports
            WHERE user_id = %s
            ORDER BY created_at DESC
            """,
            (
                session["user_id"],
            )
        )

        reports = cursor.fetchall()

        for report in reports:

            if report.get("created_at"):

                report["created_at"] = (
                    report["created_at"]
                    .strftime(
                        "%d %b %Y, %I:%M %p"
                    )
                )

        return jsonify({
            "success": True,
            "reports": reports
        })

    except Error as error:

        print(
            "Medical history error:",
            error
        )

        return jsonify({
            "success": False,
            "error":
                "Unable to load medical report history."
        }), 500

    finally:

        cursor.close()
        connection.close()


# =========================================================
# BLOOD ANALYSIS
# =========================================================

@app.route(
    "/analyze-blood",
    methods=["POST"]
)
@login_required
def analyze_blood():

    try:

        data = request.get_json(
            silent=True
        ) or {}

        fields = {

            "hemoglobin":
                str(
                    data.get(
                        "hemoglobin",
                        ""
                    )
                ).strip(),

            "blood_glucose":
                str(
                    data.get(
                        "blood_glucose",
                        ""
                    )
                ).strip(),

            "wbc":
                str(
                    data.get(
                        "wbc",
                        ""
                    )
                ).strip(),

            "platelets":
                str(
                    data.get(
                        "platelets",
                        ""
                    )
                ).strip(),

            "rbc":
                str(
                    data.get(
                        "rbc",
                        ""
                    )
                ).strip(),

            "cholesterol":
                str(
                    data.get(
                        "cholesterol",
                        ""
                    )
                ).strip(),

            "blood_pressure":
                str(
                    data.get(
                        "blood_pressure",
                        ""
                    )
                ).strip(),

            "other_values":
                str(
                    data.get(
                        "other_values",
                        ""
                    )
                ).strip()
        }

        if not any(fields.values()):

            return jsonify({
                "success": False,
                "error":
                    "Please enter at least one blood value."
            }), 400

        prompt = f"""

Interpret the following blood-test information
for general educational purposes.

Hemoglobin:
{fields["hemoglobin"] or "Not provided"}

Blood Glucose:
{fields["blood_glucose"] or "Not provided"}

WBC:
{fields["wbc"] or "Not provided"}

Platelets:
{fields["platelets"] or "Not provided"}

RBC:
{fields["rbc"] or "Not provided"}

Cholesterol:
{fields["cholesterol"] or "Not provided"}

Blood Pressure:
{fields["blood_pressure"] or "Not provided"}

Other Values:
{fields["other_values"] or "Not provided"}

Explain:

1. What each supplied value generally represents.
2. Whether the value appears within, above or below
   a reference range only when enough information exists.
3. What unusual values can sometimes be associated with.
4. Which results should be discussed with a healthcare professional.
5. Give a short overall summary.

Important:

- Reference ranges vary.
- If the user provides a reference range, prioritize it.
- Do not diagnose.
- Do not invent missing values.
- Use only values supplied by the user.

"""

        analysis = ask_huggingface(
            HEALTH_SYSTEM_PROMPT,
            prompt,
            max_tokens=1000
        )

        connection = get_db_connection()

        if connection is None:

            return jsonify({
                "success": False,
                "error":
                    "Database connection failed."
            }), 500

        cursor = connection.cursor()

        try:

            cursor.execute(
                """
                INSERT INTO blood_analysis
                (
                    user_id,
                    hemoglobin,
                    blood_glucose,
                    wbc,
                    platelets,
                    rbc,
                    cholesterol,
                    blood_pressure,
                    other_values,
                    ai_analysis
                )
                VALUES
                (
                    %s, %s, %s, %s, %s,
                    %s, %s, %s, %s, %s
                )
                """,
                (
                    session["user_id"],
                    fields["hemoglobin"],
                    fields["blood_glucose"],
                    fields["wbc"],
                    fields["platelets"],
                    fields["rbc"],
                    fields["cholesterol"],
                    fields["blood_pressure"],
                    fields["other_values"],
                    analysis
                )
            )

            connection.commit()

        finally:

            cursor.close()
            connection.close()

        return jsonify({
            "success": True,
            "analysis": analysis
        })

    except Exception as error:

        print(
            "❌ Blood Analysis Error:",
            repr(error)
        )

        return jsonify({
            "success": False,
            "error":
                "Unable to analyze blood results: "
                + str(error)
        }), 500


# =========================================================
# BLOOD HISTORY
# =========================================================

@app.route(
    "/blood-analysis-history",
    methods=["GET"]
)
@login_required
def blood_analysis_history():

    connection = get_db_connection()

    if connection is None:

        return jsonify({
            "success": False,
            "error":
                "Database connection failed."
        }), 500

    cursor = connection.cursor(
        dictionary=True
    )

    try:

        cursor.execute(
            """
            SELECT
                id,
                hemoglobin,
                blood_glucose,
                wbc,
                platelets,
                rbc,
                cholesterol,
                blood_pressure,
                other_values,
                ai_analysis,
                created_at
            FROM blood_analysis
            WHERE user_id = %s
            ORDER BY created_at DESC
            """,
            (
                session["user_id"],
            )
        )

        records = cursor.fetchall()

        for record in records:

            if record.get("created_at"):

                record["created_at"] = (
                    record["created_at"]
                    .strftime(
                        "%d %b %Y, %I:%M %p"
                    )
                )

        return jsonify({
            "success": True,
            "records": records
        })

    except Error as error:

        print(
            "Blood history error:",
            error
        )

        return jsonify({
            "success": False,
            "error":
                "Unable to load blood analysis history."
        }), 500

    finally:

        cursor.close()
        connection.close()


# =========================================================
# SYMPTOM CHECKER
# =========================================================

@app.route(
    "/check-symptoms",
    methods=["POST"]
)
@login_required
def check_symptoms():

    try:

        data = request.get_json(
            silent=True
        ) or {}

        symptoms = str(
            data.get(
                "symptoms",
                ""
            )
        ).strip()

        duration = str(
            data.get(
                "duration",
                ""
            )
        ).strip()

        severity = str(
            data.get(
                "severity",
                ""
            )
        ).strip()

        if not symptoms:

            return jsonify({
                "success": False,
                "error":
                    "Please enter your symptoms."
            }), 400

        prompt = f"""

Provide general educational guidance about these symptoms.

Symptoms:
{symptoms}

Duration:
{duration or "Not provided"}

Severity selected by user:
{severity or "Not provided"}

Explain:

1. Common possibilities that can sometimes cause these symptoms.
2. Additional information that may be useful.
3. General self-care measures when appropriate.
4. Warning signs requiring prompt or urgent medical attention.
5. When the user should consider contacting a healthcare professional.

Do NOT diagnose.
Do NOT claim certainty.
Do NOT invent symptoms or test results.

"""

        analysis = ask_huggingface(
            HEALTH_SYSTEM_PROMPT,
            prompt,
            max_tokens=900
        )

        connection = get_db_connection()

        if connection is None:

            return jsonify({
                "success": False,
                "error":
                    "Database connection failed."
            }), 500

        cursor = connection.cursor()

        try:

            cursor.execute(
                """
                INSERT INTO symptom_checks
                (
                    user_id,
                    symptoms,
                    duration,
                    severity,
                    ai_analysis
                )
                VALUES (%s, %s, %s, %s, %s)
                """,
                (
                    session["user_id"],
                    symptoms,
                    duration,
                    severity,
                    analysis
                )
            )

            connection.commit()

        finally:

            cursor.close()
            connection.close()

        return jsonify({
            "success": True,
            "analysis": analysis
        })

    except Exception as error:

        print(
            "❌ Symptom Checker Error:",
            repr(error)
        )

        return jsonify({
            "success": False,
            "error":
                "Unable to check symptoms: "
                + str(error)
        }), 500


# =========================================================
# SYMPTOM HISTORY
# =========================================================

@app.route(
    "/symptom-history",
    methods=["GET"]
)
@login_required
def symptom_history():

    connection = get_db_connection()

    if connection is None:

        return jsonify({
            "success": False,
            "error":
                "Database connection failed."
        }), 500

    cursor = connection.cursor(
        dictionary=True
    )

    try:

        cursor.execute(
            """
            SELECT
                id,
                symptoms,
                duration,
                severity,
                ai_analysis,
                created_at
            FROM symptom_checks
            WHERE user_id = %s
            ORDER BY created_at DESC
            """,
            (
                session["user_id"],
            )
        )

        records = cursor.fetchall()

        for record in records:

            if record.get("created_at"):

                record["created_at"] = (
                    record["created_at"]
                    .strftime(
                        "%d %b %Y, %I:%M %p"
                    )
                )

        return jsonify({
            "success": True,
            "records": records
        })

    except Error as error:

        print(
            "Symptom history error:",
            error
        )

        return jsonify({
            "success": False,
            "error":
                "Unable to load symptom history."
        }), 500

    finally:

        cursor.close()
        connection.close()


# =========================================================
# HEALTH INSIGHTS
# =========================================================

@app.route(
    "/generate-health-insights",
    methods=["POST"]
)
@login_required
def generate_health_insights():

    user_id = session["user_id"]

    connection = get_db_connection()

    if connection is None:

        return jsonify({
            "success": False,
            "error":
                "Database connection failed."
        }), 500

    cursor = connection.cursor(
        dictionary=True
    )

    try:

        # REPORTS
        cursor.execute(
            """
            SELECT
                report_name,
                ai_analysis,
                created_at
            FROM medical_reports
            WHERE user_id = %s
            ORDER BY created_at DESC
            LIMIT 5
            """,
            (user_id,)
        )

        reports = cursor.fetchall()

        # BLOOD
        cursor.execute(
            """
            SELECT
                hemoglobin,
                blood_glucose,
                wbc,
                platelets,
                rbc,
                cholesterol,
                blood_pressure,
                other_values,
                ai_analysis,
                created_at
            FROM blood_analysis
            WHERE user_id = %s
            ORDER BY created_at DESC
            LIMIT 5
            """,
            (user_id,)
        )

        blood_records = cursor.fetchall()

        # SYMPTOMS
        cursor.execute(
            """
            SELECT
                symptoms,
                duration,
                severity,
                ai_analysis,
                created_at
            FROM symptom_checks
            WHERE user_id = %s
            ORDER BY created_at DESC
            LIMIT 5
            """,
            (user_id,)
        )

        symptom_records = cursor.fetchall()

        # PERSONAL MEMORY
        personal_context = build_personal_context(
            user_id
        )

        context_parts = [
            "PERSONAL CONTEXT:\n"
            + personal_context
        ]

        if reports:

            context_parts.append(
                "\nMEDICAL REPORT HISTORY:"
            )

            for report in reports:

                context_parts.append(
                    f"""
Report:
{report.get("report_name", "")}

AI Analysis:
{report.get("ai_analysis", "")}
"""
                )

        if blood_records:

            context_parts.append(
                "\nBLOOD ANALYSIS HISTORY:"
            )

            for record in blood_records:

                context_parts.append(
                    f"""
Hemoglobin: {record.get("hemoglobin", "")}
Glucose: {record.get("blood_glucose", "")}
WBC: {record.get("wbc", "")}
Platelets: {record.get("platelets", "")}
RBC: {record.get("rbc", "")}
Cholesterol: {record.get("cholesterol", "")}
Blood Pressure: {record.get("blood_pressure", "")}
Other: {record.get("other_values", "")}

AI Analysis:
{record.get("ai_analysis", "")}
"""
                )

        if symptom_records:

            context_parts.append(
                "\nSYMPTOM HISTORY:"
            )

            for record in symptom_records:

                context_parts.append(
                    f"""
Symptoms:
{record.get("symptoms", "")}

Duration:
{record.get("duration", "")}

Severity:
{record.get("severity", "")}

AI Analysis:
{record.get("ai_analysis", "")}
"""
                )

        if (
            not reports
            and not blood_records
            and not symptom_records
        ):

            context_parts.append(
                "No previous medical, blood or symptom data is available."
            )

        health_context = "\n".join(
            context_parts
        )

        prompt = f"""

Create a personalized but general educational
health insight summary using the supplied saved data.

{health_context}

Structure the response as:

1. Overall Health Summary
2. Important Patterns
3. Blood/Test Observations
4. Symptom Observations
5. Positive Health Habits
6. Areas to Discuss With a Doctor
7. General Wellness Suggestions
8. Important Warning Signs, if any

IMPORTANT:

- Do not diagnose.
- Do not invent information.
- Only use information present in supplied data.
- Do not say a normal single test proves overall health.
- Do not say a normal FSH value proves fertility.
- Mention when additional clinical context is required.

"""

        insight = ask_huggingface(
            HEALTH_SYSTEM_PROMPT,
            prompt,
            max_tokens=1200
        )

        cursor.execute(
            """
            INSERT INTO health_insights
            (
                user_id,
                insight
            )
            VALUES (%s, %s)
            """,
            (
                user_id,
                insight
            )
        )

        connection.commit()

        return jsonify({
            "success": True,
            "insight": insight
        })

    except Exception as error:

        print(
            "❌ Health Insights Error:",
            repr(error)
        )

        connection.rollback()

        return jsonify({
            "success": False,
            "error":
                "Unable to generate health insights: "
                + str(error)
        }), 500

    finally:

        cursor.close()
        connection.close()


# =========================================================
# HEALTH INSIGHTS HISTORY
# =========================================================

@app.route(
    "/health-insights-history",
    methods=["GET"]
)
@login_required
def health_insights_history():

    connection = get_db_connection()

    if connection is None:

        return jsonify({
            "success": False,
            "error":
                "Database connection failed."
        }), 500

    cursor = connection.cursor(
        dictionary=True
    )

    try:

        cursor.execute(
            """
            SELECT
                id,
                insight,
                created_at
            FROM health_insights
            WHERE user_id = %s
            ORDER BY created_at DESC
            """,
            (
                session["user_id"],
            )
        )

        insights = cursor.fetchall()

        for item in insights:

            if item.get("created_at"):

                item["created_at"] = (
                    item["created_at"]
                    .strftime(
                        "%d %b %Y, %I:%M %p"
                    )
                )

        return jsonify({
            "success": True,
            "insights": insights
        })

    except Error as error:

        print(
            "Health insight history error:",
            error
        )

        return jsonify({
            "success": False,
            "error":
                "Unable to load health insight history."
        }), 500

    finally:

        cursor.close()
        connection.close()


# =========================================================
# HEALTH CHECK
# =========================================================

@app.route(
    "/api/health",
    methods=["GET"]
)
def api_health():

    db_status = False

    connection = get_db_connection()

    if connection:

        db_status = connection.is_connected()

        connection.close()

    return jsonify({

        "success": True,

        "flask": "running",

        "mysql": db_status,

        "huggingface": bool(HF_TOKEN),

        "model": HF_MODEL

    })


# =========================================================
# 404 HANDLER
# =========================================================

@app.errorhandler(404)
def page_not_found(error):

    api_paths = (

        "/ask-ai",
        "/analyze-",
        "/check-",
        "/generate-",
        "/medical-report-history",
        "/blood-analysis-history",
        "/symptom-history",
        "/health-insights-history",
        "/clear-ai-memory",
        "/memory",
        "/chat-history"

    )

    if request.path.startswith(api_paths):

        return jsonify({

            "success": False,

            "error":
                "API endpoint not found: "
                + request.path

        }), 404

    return (
        "<h1>404 - Page Not Found</h1>",
        404
    )


# =========================================================
# GLOBAL ERROR HANDLER
# =========================================================

@app.errorhandler(500)
def internal_server_error(error):

    print(
        "❌ Internal Server Error:",
        error
    )

    api_paths = (

        "/ask-ai",
        "/analyze-",
        "/check-",
        "/generate-",
        "/medical-report-history",
        "/blood-analysis-history",
        "/symptom-history",
        "/health-insights-history",
        "/clear-ai-memory",
        "/memory",
        "/chat-history"

    )

    if request.path.startswith(api_paths):

        return jsonify({

            "success": False,

            "error":
                "Internal server error. "
                "Check the Flask terminal."

        }), 500

    return (
        "<h1>500 - Internal Server Error</h1>",
        500
    )


# =========================================================
# RUN APPLICATION
# =========================================================

if __name__ == "__main__":

    print(
        "\n=========================================="
    )

    print(
        "🏥 HealthCare AI"
    )

    print(
        "=========================================="
    )

    print(
        "Database:",
        DB_NAME
    )

    print(
        "HF Model:",
        HF_MODEL
    )

    print(
        "HF Token:",
        "Configured"
        if HF_TOKEN
        else "NOT CONFIGURED"
    )

    # -----------------------------------------------------
    # INITIALIZE DATABASE TABLES
    # -----------------------------------------------------

    initialize_database()

    # -----------------------------------------------------
    # TEST DATABASE
    # -----------------------------------------------------

    connection = get_db_connection()

    if connection:

        print(
            "✅ MySQL Connected Successfully!"
        )

        connection.close()

    else:

        print(
            "❌ MySQL Connection Failed!"
        )

    print(
        "==========================================\n"
    )

    app.run(
        debug=True,
        host="127.0.0.1",
        port=5000
    )